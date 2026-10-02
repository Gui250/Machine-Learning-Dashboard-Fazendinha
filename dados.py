"""Carga de dados do dashboard.

Protocolos de inconsistência: lidos da planilha real em docs/.
Demais bases: CSVs em data/, gravados pela importação de planilhas. Base sem arquivo fica vazia.
"""
import re
from io import BytesIO
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

RAIZ = Path(__file__).parent
PASTA_DADOS = RAIZ / "data"
PLANILHA_PROTOCOLOS = RAIZ / "docs" / "Protocolo Avulso de inconsistência.xlsx"
PASTA_PROTOCOLOS = PASTA_DADOS / "protocolos"  # protocolos importados pelo painel

# Metas usadas nos gauges e comparativos. Ajuste aqui quando a diretoria definir as oficiais.
METAS = {
    "revpah": 60.0,          # R$ por hora-suíte disponível
    "ocupacao": 60.0,        # % das horas-suíte vendidas
    "conformidade": 95.0,    # % (C / (C + NC))
    "cmv": 28.0,             # % da receita de A&B (quanto menor, melhor)
    "precisao_caixa": 98.0,  # % de fechamentos às cegas sem divergência
    "deteccao_horas": 12.0,  # divergência detectada em até 12 h
}

TOTAL_SUITES = 40
CENTROS = {
    "receita_hospedagem": "Hospedagem",
    "receita_gastronomia": "Gastronomia",
    "receita_bebidas": "Bebidas & Clube do Whisky",
    "receita_entretenimento": "Entretenimento",
}

# departamento: (gestor, prioridade, turno, [(nº POP, nome, frequência)])
# Fonte: Projeto Controladoria Etapa 2 (POPs mapeados por departamento).
DEPARTAMENTOS = {
    "Departamento Pessoal": ("Felipe Steffen", "P2", None, [
        (1, "Controle de Benefícios e Vale-Transporte", "M"), (2, "Contrato para Freelancer", "M"),
        (3, "Controle de Afastamentos e Atestados", "Q"), (4, "Auditoria de Compensação de Horas", "M"),
        (5, "Conferência de Férias Programadas", "M"), (6, "Apontamentos de Rescisões", "M"),
        (7, "Apontamentos de Folha de Pagamento", "M")]),
    "Recursos Humanos": ("Felipe Steffen", "P2", None, [
        (8, "Processo de Integração (Onboarding)", "S"), (9, "Auditoria de Treinamentos Realizados", "M")]),
    "Compras": ("Felipe Steffen", "P3", None, [
        (10, "Política de 3 Cotações para Suprimentos Críticos", "S"),
        (11, "Homologação e Avaliação de Fornecedores (SLA/SLM)", "M"),
        (12, "Controle de Pedidos Urgentes (Emergenciais)", "D"), (13, "Auditoria de Nota Fiscal x Pedido de Compra", "D"),
        (14, "Conciliação de Compras por Centro de Custo", "D"), (15, "Controle de Contratos de Fornecedores", "M"),
        (16, "Conferência de Ordens de Compra", "D")]),
    "Almoxarifado": ("Felipe Steffen", "P3", None, [
        (17, "Inventário Geral de Insumos Mensal", "M"), (18, "Controle de Validade (PVPS)", "S"),
        (19, "Conciliação de Requisição x Saída de Estoque", "D"),
        (20, "Estoque de Segurança e Ponto de Ressuprimento", "Q"), (21, "Controle de Produtos Sem Giro", "M"),
        (23, "Auditoria de Requisições", "D")]),
    "Governança": ("Christiane", "P3", None, [
        (24, "Controle de Perda de Enxoval (Lavanderia)", "M"), (25, "Conciliação de Peças Recebidas vs. Devolvidas", "Q"),
        (26, "Registro de Peças Danificadas e Descarte", "S"), (27, "Conferência das Solicitações para Manutenção", "S")]),
    "Manutenção": ("Felipe Gripa", "P3", None, [
        (28, "Registro e Acompanhamento de Chamados Técnicos", "D"), (29, "Controle de Peças e Insumos de Manutenção", "M"),
        (30, "Auditoria de Ativo Imobilizado", "M")]),
    "Entretenimento": ("Maria Jeane", "P3", None, [
        (31, "Planejamento Orçamentário", "M"), (32, "SLA e SLM – Pagamentos a Prestadores", "M"),
        (33, "Gestão de Contas em Aberto", "M")]),
    "Caixa Dia": ("Carlos", "P1", "Dia", [
        (34, "Abertura e Conferência de Caixa", "D"), (35, "Conciliação de Vendas por Forma de Pagamento", "D"),
        (36, "Controle de Comandas Abertas e Fechadas", "D"), (37, "Controle de Vouchers e Cortesias", "D"),
        (38, "Controle de Descontos e Ajustes", "D"), (39, "Conciliação de Gorjetas e Taxa de Serviço", "D"),
        (40, "Fechamento Geral do Turno às Cegas", "D")]),
    "American Bar Dia": ("Carlos", "P3", "Dia", [
        (41, "Conciliação de Estoque de Bebidas", "S"), (42, "Controle de Perdas", "D"), (43, "Conferência de Consumo", "D"),
        (44, "Controle de Requisições de Produtos e Insumos", "D"), (45, "Conferência e Recebimento de Produtos", "D")]),
    "Caixa Noite": ("Ernesto", "P1", "Noite", [
        (46, "Abertura e Conferência de Caixa", "D"), (47, "Conciliação de Vendas por Forma de Pagamento", "D"),
        (48, "Controle de Comandas Abertas e Fechadas", "D"), (49, "Controle de Vouchers e Cortesias", "D"),
        (50, "Controle de Descontos e Ajustes", "D"), (51, "Conciliação de Gorjetas e Taxa de Serviço", "D"),
        (52, "Fechamento Geral do Turno às Cegas", "D")]),
    "American Bar Noite": ("Ernesto", "P3", "Noite", [
        (53, "Conciliação de Estoque de Bebidas", "S"), (54, "Controle de Perdas", "D"), (55, "Conferência de Consumo", "D"),
        (56, "Controle de Requisições de Produtos e Insumos", "D"), (57, "Conferência e Recebimento de Produtos", "D")]),
    "Recepção 360° Tarde": ("Carlos", "P1", "Dia", [
        (58, "Registro de Entrada e Saída de Clientes", "D"), (59, "Conciliação de Entrada x Caixa", "D"),
        (60, "Registro de Ocorrências de Segurança", "D")]),
    "Recepção 360° Noite": ("Christian", "P1", "Noite", [
        (61, "Registro de Entrada e Saída de Clientes", "D"), (62, "Conciliação de Entrada x Caixa", "D"),
        (63, "Registro de Ocorrências de Segurança", "D")]),
    "Cozinha Dia": ("Lilian", "P3", "Dia", [
        (64, "Aferição de Ficha Técnica e Porcionamento", "S"), (65, "Registro de Desperdício e Sobras", "D")]),
    "Cozinha Noite": ("Welika", "P3", "Noite", [
        (66, "Aferição de Ficha Técnica e Porcionamento", "S"), (67, "Controle de Produção por Previsão de Ocupação", "D")]),
}
FREQUENCIAS = {"D": "Diária", "S": "Semanal", "Q": "Quinzenal", "M": "Mensal"}
POPS_P4 = {30}  # patrimônio / ativo imobilizado: prioridade programável


def turno_da_hora(hora):
    """Dia: 06h–17h59 (Caixa Dia). Noite: 18h–05h59 (Caixa Noite)."""
    return np.where((hora >= 6) & (hora < 18), "Dia", "Noite")


def indice_conformidade(c, nc):
    """Itens Conforme ÷ (Conforme + Não Conforme) × 100. NA não entra no cálculo."""
    c, nc = float(np.sum(c)), float(np.sum(nc))
    return c / (c + nc) * 100 if c + nc else float("nan")


# ---------------------------------------------------------------- dados sintéticos (só para o self-check)

def _gerar_demo(inicio="2024-01-01", fim="2026-09-30", semente=7):
    """Bases sintéticas coerentes entre si, só para o painel funcionar antes das exportações reais."""
    rng = np.random.default_rng(semente)
    dias = pd.date_range(inicio, fim, freq="D")
    nd = len(dias)
    anos = ((dias - dias[0]).days / 365.25).to_numpy()

    # ---- receitas por hora
    sazonal_mes = np.array([1.05, 1.10, 0.98, 0.96, 0.97, 0.92, 0.94, 0.95, 0.97, 1.00, 1.03, 1.12])
    sazonal_semana = np.array([0.82, 0.84, 0.88, 0.95, 1.18, 1.25, 1.02])
    curva_hora = np.array([.74, .66, .55, .42, .30, .20, .14, .14, .16, .20, .26, .32,
                           .36, .40, .42, .44, .47, .52, .60, .68, .76, .82, .84, .80])
    fator_dia = sazonal_mes[dias.month - 1] * sazonal_semana[dias.dayofweek] * (1 + 0.05 * anos)
    paradas = rng.poisson(1.1, nd)  # suítes paradas por manutenção no dia
    hora = np.tile(np.arange(24), nd)
    d = np.repeat(np.arange(nd), 24)
    ocup = np.clip(curva_hora[hora] * fator_dia[d] * rng.normal(1, 0.10, nd * 24), 0, 0.98)
    disponiveis = TOTAL_SUITES - paradas[d]
    vendidas = np.round(ocup * disponiveis).astype(int)
    noite = turno_da_hora(hora) == "Noite"
    fim_semana = np.isin(dias.dayofweek.to_numpy()[d], [4, 5])
    preco = 92 * (1 + 0.045 * anos[d]) * np.where(noite, 1.15, 1.0) * np.where(fim_semana, 1.06, 1.0)
    receitas = pd.DataFrame({
        "data": dias[d], "hora": hora,
        "horas_disponiveis": disponiveis, "horas_vendidas": vendidas,
        "receita_hospedagem": (vendidas * preco).round(2),
        "receita_gastronomia": (vendidas * 34 * rng.lognormal(0, 0.18, nd * 24)).round(2),
        "receita_bebidas": (vendidas * np.where(noite, 52, 30) * rng.lognormal(0, 0.22, nd * 24)).round(2),
        "receita_entretenimento": np.where(np.isin(hora, [21, 22, 23, 0, 1, 2]),
                                           rng.gamma(2, 260, nd * 24) * fator_dia[d], 0).round(2),
        "comandas": rng.poisson(vendidas * 0.85),
    })

    # ---- fechamento de caixa às cegas (por dia e turno)
    receitas["turno"] = turno_da_hora(receitas["hora"])
    caixa = (receitas.assign(total=receitas[list(CENTROS)].sum(axis=1))
             .groupby(["data", "turno"], as_index=False)["total"].sum().rename(columns={"total": "valor_sistema"}))
    melhora = np.interp(((caixa["data"] - caixa["data"].min()).dt.days).to_numpy(), [0, nd], [0.16, 0.05])
    erro = np.where(rng.random(len(caixa)) < melhora, rng.normal(0, 180, len(caixa)), rng.normal(0, 1.5, len(caixa)))
    caixa["valor_contado"] = (caixa["valor_sistema"] + erro).round(2)
    receitas = receitas.drop(columns="turno")

    # ---- CMV mensal por centro
    mensal = receitas.groupby(receitas["data"].dt.to_period("M").dt.to_timestamp())[list(CENTROS)].sum()
    linhas = []
    for centro, col, teorico in [("Gastronomia", "receita_gastronomia", 0.27), ("Bebidas & Clube do Whisky", "receita_bebidas", 0.22)]:
        r = mensal[col].to_numpy()
        cmv_t = r * teorico
        cmv_r = cmv_t * (1 + rng.normal(0.05, 0.03, len(r)) * np.linspace(1.2, 0.6, len(r)))
        linhas.append(pd.DataFrame({"mes": mensal.index, "centro": centro, "receita": r.round(2),
                                    "cmv_teorico": cmv_t.round(2), "cmv_real": cmv_r.round(2),
                                    "perdas": (r * rng.normal(0.012, 0.004, len(r)).clip(0.002)).round(2)}))
    custos = pd.concat(linhas, ignore_index=True)

    # ---- orçamento mensal por centro de receita
    orcamento = (mensal.rename(columns=CENTROS).reset_index(names="mes")
                 .melt(id_vars="mes", var_name="centro", value_name="realizado"))
    orcamento["orcado"] = (orcamento["realizado"] * rng.normal(1.035, 0.045, len(orcamento))).round(2)
    orcamento = orcamento.drop(columns="realizado")

    # ---- execuções de auditoria por POP
    regras = {"D": dias, "S": dias[dias.dayofweek == 0], "Q": dias[np.isin(dias.day, [1, 16])], "M": dias[dias.day == 5]}
    taxa_nc = {"P1": 0.075, "P2": 0.045, "P3": 0.055, "P4": 0.035}
    blocos = []
    for depto, (gestor, prio, _, pops) in DEPARTAMENTOS.items():
        for num, nome, freq in pops:
            datas = regras[freq]
            n = len(datas)
            itens = 10 + num % 4
            na = rng.binomial(itens, 0.07, n)
            p = taxa_nc[prio] * np.linspace(1.6, 0.55, n) * rng.lognormal(0, 0.25, n)
            nc = rng.binomial(itens - na, np.clip(p, 0, 0.9))
            blocos.append(pd.DataFrame({
                "data": datas, "departamento": depto, "gestor": gestor,
                "prioridade": "P4" if num in POPS_P4 else prio,
                "pop_num": num, "pop": nome, "frequencia": FREQUENCIAS[freq],
                "conforme": itens - na - nc, "nao_conforme": nc, "nao_aplicavel": na,
                "horas_ate_deteccao": np.where(nc > 0, rng.gamma(2, np.linspace(6, 2.5, n)).round(1), np.nan),
            }))
    auditorias = pd.concat(blocos, ignore_index=True)

    return {"receitas": receitas, "caixa": caixa, "custos": custos, "orcamento": orcamento, "auditorias": auditorias}


BASES = ("receitas", "caixa", "custos", "orcamento", "auditorias")
COLUNAS_DATA = {"receitas": "data", "caixa": "data", "custos": "mes", "orcamento": "mes", "auditorias": "data"}


def carregar():
    """Lê data/*.csv. Base sem arquivo vem vazia, com as colunas do modelo. Retorna (bases, ausentes)."""
    bases, ausentes = {}, []
    for b in BASES:
        caminho = PASTA_DADOS / f"{b}.csv"
        if caminho.exists():
            bases[b] = pd.read_csv(caminho, parse_dates=[COLUNAS_DATA[b]])
        else:
            ausentes.append(b)
            bases[b] = pd.DataFrame({c: pd.Series(dtype="datetime64[ns]" if c == COLUNAS_DATA[b] else
                                                   object if c in TEXTO else float) for c in MODELOS[b][1]})
    bases["protocolos"] = ler_todos_protocolos()
    return bases, ausentes


# ---------------------------------------------------------------- importação por planilha modelo

# base: (título, {coluna: descrição}). A ordem das colunas é a do modelo.
MODELOS = {
    "receitas": ("Receitas por hora", {
        "data": "Data do movimento (dd/mm/aaaa)",
        "hora": "Hora do dia, de 0 a 23",
        "horas_disponiveis": "Suítes disponíveis na hora (total menos as paradas)",
        "horas_vendidas": "Suítes ocupadas na hora",
        "receita_hospedagem": "Receita de hospedagem (R$)",
        "receita_gastronomia": "Receita de gastronomia (R$)",
        "receita_bebidas": "Receita de bebidas e Clube do Whisky (R$)",
        "receita_entretenimento": "Receita de entretenimento (R$)",
        "comandas": "Número de comandas"}),
    "caixa": ("Fechamento de caixa às cegas", {
        "data": "Data do fechamento (dd/mm/aaaa)",
        "turno": "Dia ou Noite",
        "valor_sistema": "Valor apurado no Desbravador (R$)",
        "valor_contado": "Valor contado às cegas (R$)"}),
    "custos": ("Custos e CMV", {
        "mes": "Mês de referência (qualquer dia do mês)",
        "centro": "Gastronomia ou Bebidas & Clube do Whisky",
        "receita": "Receita do centro no mês (R$)",
        "cmv_teorico": "CMV pela ficha técnica (R$)",
        "cmv_real": "CMV apurado (R$)",
        "perdas": "Perdas registradas (R$)"}),
    "orcamento": ("Orçamento", {
        "mes": "Mês de referência (qualquer dia do mês)",
        "centro": ", ".join(CENTROS.values()),
        "orcado": "Receita orçada (R$)"}),
    "auditorias": ("Auditorias POP", {
        "data": "Data da auditoria (dd/mm/aaaa)",
        "departamento": "Departamento auditado",
        "gestor": "Gestor do departamento",
        "prioridade": "P1, P2, P3 ou P4",
        "pop_num": "Número do POP",
        "pop": "Nome do POP",
        "frequencia": ", ".join(FREQUENCIAS.values()),
        "conforme": "Itens conformes (C)",
        "nao_conforme": "Itens não conformes (NC)",
        "nao_aplicavel": "Itens não aplicáveis (NA)",
        "horas_ate_deteccao": "Opcional: horas até detectar a divergência"}),
}
CHAVES = {"receitas": ["data", "hora"], "caixa": ["data", "turno"], "custos": ["mes", "centro"],
          "orcamento": ["mes", "centro"], "auditorias": ["data", "pop_num"]}
TEXTO = {"turno", "centro", "departamento", "gestor", "prioridade", "pop", "frequencia"}
OPCIONAIS = {"horas_ate_deteccao"}
PERMITIDOS = {
    ("caixa", "turno"): {"Dia", "Noite"},
    ("custos", "centro"): {"Gastronomia", "Bebidas & Clube do Whisky"},
    ("orcamento", "centro"): set(CENTROS.values()),
    ("auditorias", "prioridade"): {"P1", "P2", "P3", "P4"},
    ("auditorias", "frequencia"): set(FREQUENCIAS.values()),
}


def _numero(s):
    """Aceita números do Excel e texto no formato brasileiro ("R$ 1.234,56")."""
    if pd.api.types.is_numeric_dtype(s):
        return s
    t = s.astype(str).str.replace("R$", "", regex=False).str.strip()
    virgula = t.str.contains(",", regex=False)
    t = t.where(~virgula, t.str.replace(".", "", regex=False).str.replace(",", ".", regex=False))
    return pd.to_numeric(t, errors="coerce")


def _data(s):
    """ISO (aaaa-mm-dd) primeiro; o resto é lido como dd/mm/aaaa. dayfirst sozinho inverte datas ISO."""
    iso = pd.to_datetime(s, format="ISO8601", errors="coerce")
    return iso.fillna(pd.to_datetime(s, dayfirst=True, format="mixed", errors="coerce"))


def ler_planilha(arquivo):
    """Lê o upload (.xlsx: primeira aba; .csv: separador e codificação detectados)."""
    try:
        if not arquivo.name.lower().endswith(".csv"):
            return pd.read_excel(arquivo, sheet_name=0)
        try:
            return pd.read_csv(arquivo, sep=None, engine="python", encoding="utf-8-sig")
        except UnicodeDecodeError:
            arquivo.seek(0)
            return pd.read_csv(arquivo, sep=None, engine="python", encoding="latin-1")
    except Exception as erro:
        raise ValueError(f"Não foi possível ler o arquivo: {erro}") from erro


def validar(base, df):
    """Confere a planilha contra o modelo da base. Devolve a base limpa ou levanta ValueError com os problemas."""
    colunas = list(MODELOS[base][1])
    df = df.rename(columns=lambda c: str(c).strip().lower())
    faltam = [c for c in colunas if c not in df.columns]
    if faltam:
        raise ValueError(f"Colunas ausentes: {', '.join(faltam)}. Use o modelo da base.")
    df = df[colunas].dropna(how="all").copy()
    if df.empty:
        raise ValueError("A planilha não tem linhas preenchidas.")
    problemas = []
    for c in colunas:
        bruto = df[c]
        if c == COLUNAS_DATA[base]:
            df[c] = _data(bruto)
            if c == "mes":
                df[c] = df[c].dt.to_period("M").dt.to_timestamp()
        elif c in TEXTO:
            df[c] = bruto.astype("string").str.strip()
        else:
            df[c] = _numero(bruto)
        ruins = df[c].isna() & bruto.notna() if c in OPCIONAIS else df[c].isna()
        if (base, c) in PERMITIDOS:
            ruins |= ~df[c].isin(PERMITIDOS[base, c]).fillna(False)
        if c == "hora":
            ruins |= ~df[c].between(0, 23)
        if ruins.any():
            i = ruins.idxmax()
            exemplo = "vazio" if pd.isna(bruto[i]) else f'"{bruto[i]}"'
            problemas.append(f"{c}: {ruins.sum()} linha(s) inválida(s), por exemplo a linha {i + 2} ({exemplo})")
    if problemas:
        raise ValueError("\n".join(problemas))
    return df


def importar(base, df, substituir=False):
    """Grava a base validada. Sem substituir, acrescenta e atualiza as linhas com a mesma chave (ex.: data + hora)."""
    PASTA_DADOS.mkdir(exist_ok=True)
    caminho = PASTA_DADOS / f"{base}.csv"
    if not substituir and caminho.exists():
        df = pd.concat([pd.read_csv(caminho, parse_dates=[COLUNAS_DATA[base]]), df], ignore_index=True)
    df.drop_duplicates(CHAVES[base], keep="last").sort_values(CHAVES[base]).to_csv(caminho, index=False)


def modelo_planilha(base, exemplo):
    """Modelo .xlsx: aba da base com cabeçalho e 3 linhas de exemplo, mais a aba Instruções."""
    titulo, colunas = MODELOS[base]
    exemplo = exemplo[list(colunas)].head(3).copy()
    exemplo[COLUNAS_DATA[base]] = exemplo[COLUNAS_DATA[base]].dt.date
    instrucoes = pd.DataFrame({"coluna": list(colunas), "descrição": list(colunas.values())})
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        exemplo.to_excel(w, sheet_name=base, index=False)
        instrucoes.to_excel(w, sheet_name="Instruções", index=False)
        for ws in w.book.worksheets:
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(len(str(x.value or "")) for x in col) + 3
    return buf.getvalue()


def modelo_protocolo():
    """A própria planilha de protocolos, só com a aba MODELO (formatação preservada)."""
    wb = openpyxl.load_workbook(PLANILHA_PROTOCOLOS)
    for ws in wb.worksheets:
        if ws.title.strip().upper() != "MODELO":
            wb.remove(ws)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def importar_protocolos(conteudo, nome="protocolos.xlsx"):
    """Valida e guarda uma planilha de protocolos em data/protocolos/. Mesmo número substitui o anterior."""
    try:
        p = ler_protocolos(BytesIO(conteudo))
    except Exception as erro:
        raise ValueError(f"Não foi possível ler a planilha: {erro}") from erro
    if p.empty:
        raise ValueError("Nenhum protocolo encontrado. Duplique a aba MODELO e preencha uma aba por protocolo.")
    incompletos = p.loc[(p["departamento"] == "") | (p["descricao"] == ""), "protocolo"].tolist()
    if incompletos:
        raise ValueError(f"Preencha Departamento e Descrição da Inconsistência nos protocolos: {', '.join(incompletos)}.")
    PASTA_PROTOCOLOS.mkdir(parents=True, exist_ok=True)
    (PASTA_PROTOCOLOS / f"{pd.Timestamp.now():%Y%m%d-%H%M%S}-{Path(nome).name}").write_bytes(conteudo)
    return p


def ler_todos_protocolos():
    """Planilha original em docs/ mais as importadas, na ordem; protocolo repetido fica com a versão mais nova."""
    arquivos = [PLANILHA_PROTOCOLOS, *sorted(PASTA_PROTOCOLOS.glob("*.xlsx"))]
    return (pd.concat([ler_protocolos(a) for a in arquivos], ignore_index=True)
            .drop_duplicates("protocolo", keep="last").reset_index(drop=True))

_SECOES = ("Descrição dos POPs", "Descrição da Inconsistência", "Devolutiva da Diretoria", "2.  PROTOCOLO PDCA")


def _texto(v):
    return str(v).strip() if v is not None else ""


def ler_protocolos(caminho=PLANILHA_PROTOCOLOS) -> pd.DataFrame:
    """Um registro por aba (exceto MODELO) da planilha de Protocolo Avulso de Inconsistência."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    registros = []
    for ws in wb.worksheets:
        if ws.title.strip().upper() == "MODELO":
            continue
        campos, secoes, pdca, secao = {}, {s: [] for s in _SECOES[:3]}, {}, None
        for linha in ws.iter_rows(values_only=True):
            celulas = [_texto(c) for c in linha]
            rotulo, resto = celulas[0], [c for c in celulas[1:] if c]
            if not rotulo and not resto:
                continue
            if rotulo in _SECOES:
                secao = rotulo if rotulo in secoes else None
                secoes.get(rotulo, []).extend(resto)
                continue
            if rotulo[:3] in ("P —", "D —", "C —", "A —"):
                pdca[rotulo[0]] = " ".join(resto)
                secao = None
                continue
            if secao:
                secoes[secao].extend([rotulo] + resto if rotulo else resto)
            elif resto:
                campos.setdefault(rotulo, resto[0])
            elif rotulo.startswith("Data"):
                campos.setdefault("Data", rotulo)

        descricao = " ".join(x for x in secoes["Descrição da Inconsistência"] if x)
        datas = re.findall(r"\d{2}/\d{2}/\d{4}", campos.get("Data", "")) or re.findall(r"\d{2}/\d{2}/\d{4}", descricao)
        devolutiva = " ".join(x for x in secoes["Devolutiva da Diretoria"] if x)
        registros.append({
            "protocolo": _texto(campos.get("Nº DO PROTOCOLO", ws.title)).zfill(3),
            "data": pd.to_datetime(datas[0], dayfirst=True) if datas else pd.NaT,
            "departamento": campos.get("Departamento", ""),
            "responsavel": campos.get("Responsável (R)", ""),
            "gestor": campos.get("Gestor", ""),
            "setor": campos.get("Setor", ""),
            "pops": [p for p in secoes["Descrição dos POPs"] if p],
            "descricao": descricao,
            "devolutiva": devolutiva,
            "status": "Respondido" if devolutiva else "Aguardando devolutiva",
            **{f"pdca_{k}": pdca.get(k, "") for k in "PDCA"},
        })
    return pd.DataFrame(registros)


if __name__ == "__main__":
    assert indice_conformidade([8], [2]) == 80.0  # NA fica fora do cálculo
    assert list(turno_da_hora(np.array([5, 6, 17, 18]))) == ["Noite", "Dia", "Dia", "Noite"]
    p = ler_protocolos()
    assert list(p["protocolo"]) == ["001", "002", "003"] and p["pops"].map(len).tolist() == [7, 3, 2]
    g = _gerar_demo(fim="2024-03-31")
    a = g["auditorias"]
    assert (a[["conforme", "nao_conforme", "nao_aplicavel"]] >= 0).all().all()
    assert a["pop_num"].nunique() == 66  # POPs 1–67 (o 22 não existe na Etapa 2)
    r = g["receitas"]
    assert (r["horas_vendidas"] <= r["horas_disponiveis"]).all()
    # importação: o modelo baixado volta íntegro; texto brasileiro e datas dd/mm são aceitos
    for b in MODELOS:
        ida = pd.read_excel(BytesIO(modelo_planilha(b, g[b])), sheet_name=0)
        pd.testing.assert_frame_equal(validar(b, ida).reset_index(drop=True), g[b][list(MODELOS[b][1])].head(3),
                                      check_dtype=False, check_exact=False)
    cx = validar("caixa", pd.DataFrame({"Data": ["05/01/2024", "2024-01-05"], "turno": ["Dia", "Noite"],
                                        "valor_sistema": ["R$ 1.234,56", "10"], "valor_contado": [1234.5, 9.5]}))
    assert list(cx["data"]) == [pd.Timestamp("2024-01-05")] * 2 and cx["valor_sistema"].tolist() == [1234.56, 10]
    for ruim in ({"turno": ["Tarde", "Dia"]}, {"valor_contado": ["abc", 1]}):
        try:
            validar("caixa", cx.assign(**ruim))
            raise AssertionError(ruim)
        except ValueError:
            pass
    assert ler_protocolos(BytesIO(modelo_protocolo())).empty
    print("ok")
